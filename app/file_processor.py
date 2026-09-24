import csv
import io
import json
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from langchain_core.documents import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation



SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".md",
    ".csv",
    ".xlsx",
    ".pptx",
    ".json",
    ".xml",
    ".html",
    ".htm",
    ".zip",
}



def extract_pdf(
    file_bytes: bytes,
    source_name: str,
) -> list[Document]:
    reader = PdfReader(
        io.BytesIO(file_bytes)
    )

    documents: list[Document] = []

    for page_number, page in enumerate(
        reader.pages,
        start=1,
    ):
        text = page.extract_text() or ""

        if not text.strip():
            continue

        documents.append(
            Document(
                page_content=text,
                metadata={
                    "source": source_name,
                    "file_type": ".pdf",
                    "page": page_number,
                },
            )
        )

    return documents



def extract_docx(
    file_bytes: bytes,
    source_name: str,
) -> list[Document]:
    document = DocxDocument(
        io.BytesIO(file_bytes)
    )

    parts: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            parts.append(text)

    # Include table contents.
    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
            ]

            if any(cells):
                parts.append(
                    " | ".join(cells)
                )

    text = "\n".join(parts)

    if not text.strip():
        return []

    return [
        Document(
            page_content=text,
            metadata={
                "source": source_name,
                "file_type": ".docx",
            },
        )
    ]



def extract_text_file(
    file_bytes: bytes,
    source_name: str,
    extension: str,
) -> list[Document]:
    text = file_bytes.decode(
        "utf-8",
        errors="replace",
    )

    if not text.strip():
        return []

    return [
        Document(
            page_content=text,
            metadata={
                "source": source_name,
                "file_type": extension,
            },
        )
    ]



def extract_csv(
    file_bytes: bytes,
    source_name: str,
) -> list[Document]:
    decoded = file_bytes.decode(
        "utf-8",
        errors="replace",
    )

    reader = csv.reader(
        io.StringIO(decoded)
    )

    rows: list[str] = []

    for row in reader:
        rows.append(
            " | ".join(row)
        )

    text = "\n".join(rows)

    if not text.strip():
        return []

    return [
        Document(
            page_content=text,
            metadata={
                "source": source_name,
                "file_type": ".csv",
            },
        )
    ]




def extract_xlsx(
    file_bytes: bytes,
    source_name: str,
) -> list[Document]:
    workbook = load_workbook(
        io.BytesIO(file_bytes),
        read_only=True,
        data_only=True,
    )

    documents: list[Document] = []

    for worksheet in workbook.worksheets:
        parts = [
            f"[Sheet: {worksheet.title}]"
        ]

        for row in worksheet.iter_rows(
            values_only=True
        ):
            values = [
                ""
                if value is None
                else str(value)
                for value in row
            ]

            if any(values):
                parts.append(
                    " | ".join(values)
                )

        text = "\n".join(parts)

        if text.strip():
            documents.append(
                Document(
                    page_content=text,
                    metadata={
                        "source": source_name,
                        "file_type": ".xlsx",
                        "sheet": worksheet.title,
                    },
                )
            )

    workbook.close()

    return documents




def extract_pptx(
    file_bytes: bytes,
    source_name: str,
) -> list[Document]:
    presentation = Presentation(
        io.BytesIO(file_bytes)
    )

    documents: list[Document] = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        parts = [
            f"[Slide {slide_number}]"
        ]

        for shape in slide.shapes:
            if not hasattr(shape, "text"):
                continue

            text = shape.text.strip()

            if text:
                parts.append(text)

        text = "\n".join(parts)

        if text.strip():
            documents.append(
                Document(
                    page_content=text,
                    metadata={
                        "source": source_name,
                        "file_type": ".pptx",
                        "slide": slide_number,
                    },
                )
            )

    return documents



def extract_json(
    file_bytes: bytes,
    source_name: str,
) -> list[Document]:
    raw_text = file_bytes.decode(
        "utf-8",
        errors="replace",
    )

    data = json.loads(raw_text)

    text = json.dumps(
        data,
        indent=2,
        ensure_ascii=False,
    )

    if not text.strip():
        return []

    return [
        Document(
            page_content=text,
            metadata={
                "source": source_name,
                "file_type": ".json",
            },
        )
    ]



def extract_xml(
    file_bytes: bytes,
    source_name: str,
) -> list[Document]:
    root = ElementTree.fromstring(
        file_bytes
    )

    parts: list[str] = []

    for element in root.iter():
        if (
            element.text
            and element.text.strip()
        ):
            parts.append(
                element.text.strip()
            )

    text = "\n".join(parts)

    if not text.strip():
        return []

    return [
        Document(
            page_content=text,
            metadata={
                "source": source_name,
                "file_type": ".xml",
            },
        )
    ]




def extract_html(
    file_bytes: bytes,
    source_name: str,
    extension: str,
) -> list[Document]:
    decoded = file_bytes.decode(
        "utf-8",
        errors="replace",
    )

    soup = BeautifulSoup(
        decoded,
        "html.parser",
    )

    text = soup.get_text(
        separator="\n",
        strip=True,
    )

    if not text.strip():
        return []

    return [
        Document(
            page_content=text,
            metadata={
                "source": source_name,
                "file_type": extension,
            },
        )
    ]




def extract_zip(
    file_bytes: bytes,
    source_name: str,
    user_id: int | None,
    file_id: int | None,
    max_files: int,
    max_uncompressed_size: int,
) -> list[Document]:
    documents: list[Document] = []

    with zipfile.ZipFile(
        io.BytesIO(file_bytes)
    ) as archive:

        entries = archive.infolist()

        if len(entries) > max_files:
            raise ValueError(
                "ZIP contains too many files."
            )

        total_uncompressed_size = sum(
            entry.file_size
            for entry in entries
            if not entry.is_dir()
        )

        if (
            total_uncompressed_size
            > max_uncompressed_size
        ):
            raise ValueError(
                "ZIP expands beyond the allowed size."
            )

        for entry in entries:

            if entry.is_dir():
                continue

            nested_name = entry.filename
            nested_extension = (
                Path(nested_name).suffix.lower()
            )

          
            if nested_extension == ".zip":
                continue

            if (
                nested_extension
                not in SUPPORTED_EXTENSIONS
            ):
                continue

            nested_bytes = archive.read(
                entry
            )

            extracted_documents = (
                extract_documents(
                    filename=nested_name,
                    file_bytes=nested_bytes,
                    user_id=user_id,
                    file_id=file_id,
                    max_zip_files=max_files,
                    max_zip_uncompressed_size=(
                        max_uncompressed_size
                    ),
                )
            )

            for document in extracted_documents:

              
                document.metadata[
                    "source"
                ] = nested_name

                document.metadata[
                    "archive"
                ] = source_name

                documents.append(
                    document
                )

    return documents




def extract_documents(
    filename: str,
    file_bytes: bytes,
    user_id: int | None = None,
    file_id: int | None = None,
    max_zip_files: int = 100,
    max_zip_uncompressed_size: int = 200 * 1024 * 1024,
) -> list[Document]:

    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(
            sorted(SUPPORTED_EXTENSIONS)
        )

        raise ValueError(
            f"Unsupported file type '{extension}'. "
            f"Supported file types: {supported}"
        )

    if extension == ".pdf":

        documents = extract_pdf(
            file_bytes,
            filename,
        )

    elif extension == ".docx":

        documents = extract_docx(
            file_bytes,
            filename,
        )

    elif extension in {".txt", ".md"}:

        documents = extract_text_file(
            file_bytes,
            filename,
            extension,
        )

    elif extension == ".csv":

        documents = extract_csv(
            file_bytes,
            filename,
        )

    elif extension == ".xlsx":

        documents = extract_xlsx(
            file_bytes,
            filename,
        )

    elif extension == ".pptx":

        documents = extract_pptx(
            file_bytes,
            filename,
        )

    elif extension == ".json":

        documents = extract_json(
            file_bytes,
            filename,
        )

    elif extension == ".xml":

        documents = extract_xml(
            file_bytes,
            filename,
        )

    elif extension in {".html", ".htm"}:

        documents = extract_html(
            file_bytes,
            filename,
            extension,
        )

    elif extension == ".zip":

        documents = extract_zip(
            file_bytes=file_bytes,
            source_name=filename,
            user_id=user_id,
            file_id=file_id,
            max_files=max_zip_files,
            max_uncompressed_size=(
                max_zip_uncompressed_size
            ),
        )

    else:
        raise ValueError(
            f"No processor available for '{extension}'."
        )

 
    for document in documents:

        if user_id is not None:
            document.metadata[
                "user_id"
            ] = str(user_id)

        if file_id is not None:
            document.metadata[
                "file_id"
            ] = str(file_id)

    return documents